# -*- coding: utf-8 -*-
import pptx
import os
import shutil

def set_shape_text(shape, text_or_lines):
    tf = shape.text_frame
    if isinstance(text_or_lines, str):
        lines = [text_or_lines]
    else:
        lines = text_or_lines

    # Remember formatting from first run if available
    font_name = None
    font_size = None
    font_bold = None
    font_color_rgb = None

    if tf.paragraphs and tf.paragraphs[0].runs:
        r0 = tf.paragraphs[0].runs[0]
        font_name = r0.font.name
        font_size = r0.font.size
        font_bold = r0.font.bold
        try:
            if r0.font.color and r0.font.color.rgb:
                font_color_rgb = r0.font.color.rgb
        except Exception:
            pass

    for i, line in enumerate(lines):
        if i < len(tf.paragraphs):
            p = tf.paragraphs[i]
            if p.runs:
                p.runs[0].text = line
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.text = line
        else:
            p = tf.add_paragraph()
            r = p.add_run()
            r.text = line
            if font_name:
                r.font.name = font_name
            if font_size:
                r.font.size = font_size
            if font_bold is not None:
                r.font.bold = font_bold
            if font_color_rgb:
                r.font.color.rgb = font_color_rgb

    # Clear any extra paragraphs
    if len(tf.paragraphs) > len(lines):
        for i in range(len(lines), len(tf.paragraphs)):
            p = tf.paragraphs[i]
            if p.runs:
                for r in p.runs:
                    r.text = ""
            else:
                p.text = ""

def update_presentation(src_path, dst_path):
    prs = pptx.Presentation(src_path)

    # ----------------------------------------------------
    # Slide 2: 핵심 전략 요약
    # ----------------------------------------------------
    s2 = prs.slides[1]
    for shape in s2.shapes:
        if shape.name == "Text 13":
            set_shape_text(shape, [
                "· 정적 카탈로그형 사이트의 한계를 넘어 인터랙티브 웹으로 전환",
                "· 고화질 비디오 배경 & 스크롤에 반응하는 모던 스토리텔링 중심 설계"
            ])
        elif shape.name == "Text 18":
            set_shape_text(shape, [
                "· 깜빡임 없이 부드러운 화면 전환(Seamless) 인터랙션 설계",
                "· 스크롤 흐름에 맞춘 부드러운 화면 제어 & 텍스트 모션"
            ])
        elif shape.name == "Text 38":
            set_shape_text(shape, [
                "· 5단계 국문/영문 심볼릭 카피라이팅 시스템",
                "· 고화질 영상, 부드러운 웹 모션, 반응형 레이아웃 기술 융합"
            ])

    # ----------------------------------------------------
    # Slide 3: 핵심 디자인 & 기술 컨셉
    # ----------------------------------------------------
    s3 = prs.slides[2]
    for shape in s3.shapes:
        if shape.name == "Text 27":
            set_shape_text(shape, "시네마틱 고화질 영상 + 감각적인 모션 연출")
        elif shape.name == "Text 29":
            set_shape_text(shape, [
                "· 미래지향적 비디오 배경(futuristic_key.mp4)과 그래픽 연출로 시각적 신뢰감 제공",
                "· 99.98% 가동률, 50+ 엔터프라이즈 레퍼런스를 직관적 수치로 증명"
            ])

    # ----------------------------------------------------
    # Slide 4: 유기적인 페이지 이동 및 네비게이팅 시스템 (3번 첨부 이미지)
    # ----------------------------------------------------
    s4 = prs.slides[3]
    for shape in s4.shapes:
        if shape.name == "Text 4":
            set_shape_text(shape, "깜빡임 없는 매끄러운 화면 전환과 풀스크린 커튼 메뉴 시스템")
        elif shape.name == "Text 12":
            set_shape_text(shape, [
                "· 화면 최상단에 항상 고정되어 언제든 원하는 메뉴로 이동 가능",
                "· 스크롤 시작 시 은은한 반투명 다크 배경으로 본문과 자연스럽게 조화",
                "· 본문 콘텐츠를 가리지 않으면서 로고 / 메뉴 버튼 구조를 깔끔하게 유지",
                "· WINHUB 4대 핵심 메뉴 [Home · About Us · Works · Contact Us] 직관적 노출"
            ])
        elif shape.name == "Text 16":
            set_shape_text(shape, [
                "· 좁은 드롭다운 창 대신 화면 전체(풀스크린)를 감각적으로 덮는 시원한 오버레이 메뉴",
                "· 메뉴 텍스트: 아래에서 위로 물 흐르듯 순차적으로 나타나는 트렌디한 모션 연출",
                "· 큼직하고 세련된 폰트를 적용하여 방문자에게 명확하고 직관적인 브랜드 인상 전달",
                "· 참조 사이트: 1367studio (https://www.1367studio.com/) 풀스크린 메뉴 스타일 벤치마킹",
                "· ESC 키 감지 및 닫기 버튼으로 언제든지 편하고 빠르게 닫을 수 있는 쉬운 사용성"
            ])
        elif shape.name == "Text 19":
            set_shape_text(shape, "3. 화면 깜빡임 없는 매끄러운 페이지 이동 (화면 끊김 방지)")
        elif shape.name == "Text 20":
            set_shape_text(shape, [
                "· 1단계: 메뉴 클릭 시 화면이 하얗게 깜빡이지 않고 부드럽게 전환 시작",
                "· 2단계: 현재 화면이 은은하게 페이드아웃되며 다음 페이지로 자연스럽게 연결",
                "· 3단계: 상단 헤더 메뉴와 배경 무드를 안정적으로 유지하며 새 콘텐츠 신속 로딩",
                "· 4단계: 새로운 페이지의 최상단 위치로 스크롤을 매끄럽게 맞추어 즉시 탐색 지원",
                "· 5단계: 신규 페이지의 인터랙티브 요소들이 유기적으로 등장하며 방문자 몰입 유도"
            ])

    # ----------------------------------------------------
    # Slide 5: 인터랙션 및 모션 시스템
    # ----------------------------------------------------
    s5 = prs.slides[4]
    for shape in s5.shapes:
        if shape.name == "Text 24":
            set_shape_text(shape, [
                "· 화면이 아래로 밀려나지 않고 고정된 채, 스크롤에 맞춰 텍스트와 그래픽이 자연스럽게 변화",
                "· 영상 배경과 그래픽 모션, 텍스트 효과가 스크롤 흐름에 맞추어 완벽하게 결합"
            ])
        elif shape.name == "Text 27":
            set_shape_text(shape, "텍스트 순차 등장 모션 (Line Reveal)")

    # ----------------------------------------------------
    # Slide 6: 기술 스택 요약 테이블
    # ----------------------------------------------------
    s6 = prs.slides[5]
    for shape in s6.shapes:
        if shape.has_table:
            table = shape.table
            for row in table.rows:
                if len(row.cells) >= 3 and "Three.js" in row.cells[1].text:
                    row.cells[1].text = "HTML5 Video & Modern SVG"
                    row.cells[2].text = "고화질 배경 비디오, 심플 벡터 그래픽, 실시간 부드러운 웹 모션"
                    if len(row.cells) >= 4:
                        row.cells[3].text = "futuristic_key 비디오 + 세련된 AI/레이더 그래픽 연출"
                elif len(row.cells) >= 3 and "Motion Orchestration" in row.cells[0].text:
                    row.cells[2].text = "스크롤 속도에 맞추어 화면과 텍스트가 자연스럽게 변화하는 1:1 모션 연동"
                elif len(row.cells) >= 3 and "Typography Engine" in row.cells[0].text:
                    row.cells[2].text = "텍스트 문장을 행 단위로 분할하여 아래에서 위로 부드럽게 나타나는 등장 효과"

    # ----------------------------------------------------
    # Slide 9: 네비게이션 상세 스펙
    # ----------------------------------------------------
    s9 = prs.slides[8]
    for shape in s9.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if "Clip-path Reveal" in p.text:
                    p.text = "▶ [텍스트 순차 등장]: 01 Home부터 04 Contact Us까지 텍스트가 아래에서 위로 물 흐르듯 순차적으로 나타남"
                elif "백드롭 블러(blur: 24px)" in p.text:
                    p.text = "▶ [진입 모션]: 메뉴 버튼 클릭 시 0.4초 만에 은은한 반투명 다크 배경과 함께 풀스크린 메뉴 창이 시원하게 열림"

    # ----------------------------------------------------
    # Slide 11: 메인 화면 - Hero 섹션 (Act 1)
    # ----------------------------------------------------
    s11 = prs.slides[10]
    for shape in s11.shapes:
        if shape.name == "Text 40":
            set_shape_text(shape, [
                "▶ [고화질 배경 영상 연동]:",
                "  · assets/mp4/futuristic_key.mp4 고화질 비디오 무한 루프 재생",
                "  · 다크 오버레이를 적용하여 메인 텍스트의 선명한 가독성 확보",
                "▶ [직관적인 메인 카피]:",
                "  · 'CONNECTS TO EMPOWER YOUR SUCCESS' 슬로건으로 비즈니스 도약 비전 제시",
                "▶ [GNB 및 화면 이동]:",
                "  · 상단 고정 로고 + 햄버거 메뉴 배치 및 아래로 부드럽게 스크롤을 유도하는 인터랙션"
            ])

    # ----------------------------------------------------
    # Slide 12: 메인 화면 - S.01 Tech Innovation (Act 2) - 10p/22p 일치 & 3D 삭제!
    # ----------------------------------------------------
    s12 = prs.slides[11]
    for shape in s12.shapes:
        if shape.name == "Text 34":
            set_shape_text(shape, "AI 칩셋 & 신경망 그래픽")
        elif shape.name == "Text 36":
            set_shape_text(shape, "은은한 무드 배경 효과")
        elif shape.name == "Text 39":
            set_shape_text(shape, [
                "▶ [AI 칩셋 & 신경망 비주얼]:",
                "  · 심플 AI 칩셋 SVG 그래픽과 데이터 흐름 라인이 스크롤에 맞춰 은은하게 반응",
                "▶ [핵심 기술 역량 전달]:",
                "  · AI, 클라우드, 빅데이터를 기업 환경에 맞춰 유기적으로 설계하고 운영하는 역량 표현",
                "▶ [반응형 레이아웃]:",
                "  · PC, 태블릿, 모바일 모든 기기에서 텍스트와 그래픽이 한눈에 들어오는 가독성 중심 디자인"
            ])

    # ----------------------------------------------------
    # Slide 13: 메인 화면 - S.02 Custom Strategy (Act 3) - 10p/22p 일치 & 3D 삭제!
    # ----------------------------------------------------
    s13 = prs.slides[12]
    for shape in s13.shapes:
        if shape.name == "Text 35":
            set_shape_text(shape, [
                "360° 레이더 스캐너",
                "[맞춤형 DX 로드맵 그래픽]"
            ])
        elif shape.name == "Text 38":
            set_shape_text(shape, [
                "▶ [정밀 레이더 스캔 모션]:",
                "  · 360도 회전하는 레이더 그래픽으로 고객 비즈니스 환경을 정밀 분석하는 컨셉 시각화",
                "▶ [맞춤 전략 로드맵]:",
                "  · 정형화된 솔루션이 아닌 기업별 요구사항에 꼭 맞춘 단계별 시스템 구축 및 운영 로드맵 제시",
                "▶ [모던 다크 스타일]:",
                "  · 군더더기 없는 얇은 선과 포인트 컬러로 정밀함과 전문성을 강조"
            ])

    # ----------------------------------------------------
    # Slide 14: 메인 화면 - S.03 Enterprise Impact (Act 4) - 10p/22p 일치 & 3D 삭제!
    # ----------------------------------------------------
    s14 = prs.slides[13]
    for shape in s14.shapes:
        if shape.name == "Text 32":
            set_shape_text(shape, "차세대 인터랙티브 2x3 수치 그리드")
        elif shape.name == "Text 53":
            set_shape_text(shape, [
                "▶ [숫자로 입증하는 신뢰 (2x3 그리드)]:",
                "  · 99.98% 가동률, 15년+ 경험, 50개+ 대기업 파트너 등 핵심 수치 카운트업 효과",
                "▶ [직관적인 정보 구조]:",
                "  · 카드 테두리 없는 모던한 얇은 구분선으로 6개 핵심 지표를 한눈에 전달",
                "▶ [검증된 레퍼런스 강조]:",
                "  · 말로만 하는 기술이 아닌 실질적인 운영 실적을 통해 고객 신뢰 확보"
            ])

    # ----------------------------------------------------
    # Slide 15: 메인 화면 - S.04 Works Showcase (Act 5 전반부) - 전문용어 순화!
    # ----------------------------------------------------
    s15 = prs.slides[14]
    for shape in s15.shapes:
        if shape.name == "TextBox 60":
            set_shape_text(shape, [
                "▶ [한눈에 보는 실적 목록]:",
                "  · ABC마트, LX하우시스, AK플라자, 신세계I&C 등 대표 프로젝트를 깔끔한 목록으로 제시",
                "  · 마우스를 올리면 각 프로젝트의 대표 이미지와 정보가 부드럽게 나타남",
                "▶ [실적 데이터 연동]:",
                "  · 주요 엔터프라이즈 레퍼런스를 직관적으로 노출하고 신뢰감 부여",
                "  · 클릭 시 Works 상세 화면으로 매끄럽게 이동하여 심층 내용 확인",
                "▶ [엔터프라이즈 최적화 디자인]:",
                "  · 얇은 구분선과 시안 컬러 포인트로 세련된 가독성 극대화",
                "  · 모바일 및 태블릿 터치 환경에 최적화된 편리한 터치 영역 보장"
            ])

    # ----------------------------------------------------
    # Slide 16: 메인 화면 - S.05 Outro CTA (Act 5 후반부) - 10p/22p 일치 & 3D 삭제!
    # ----------------------------------------------------
    s16 = prs.slides[15]
    for shape in s16.shapes:
        if shape.name == "Text 38":
            set_shape_text(shape, [
                "▶ [명확한 문의 유도 (CTA)]:",
                "  · '프로젝트 제안 문의하기' 버튼을 강조하여 실질적인 비즈니스 상담 연결",
                "▶ [하단 푸터 정보]:",
                "  · 회사 소개, 주요 메뉴 링크(About Us, Works, Contact Us), 연락처 및 본사 위치 안내",
                "▶ [완결성 있는 스토리 마무리]:",
                "  · 메인 영상(Act 1)부터 문의 제안(Act 5)까지 일관된 흐름으로 완결"
            ])

    # ----------------------------------------------------
    # Slide 17: 서브페이지 - ABOUT US (히어로) - 3D 삭제!
    # ----------------------------------------------------
    s17 = prs.slides[16]
    for shape in s17.shapes:
        if shape.name == "TextBox 48":
            set_shape_text(shape, "서브페이지 - ABOUT US (기업 비전 & 시네마틱 영상 히어로)")
        elif shape.name == "Text 44":
            set_shape_text(shape, [
                "▶ [시네마틱 영상 비주얼 연출]:",
                "  · wire_circle_network.mp4 비디오 배경과 기업 철학을 연결하여 세련된 몰입감 제공",
                "▶ [기업 비전]:",
                "  · 'WINHUB는 비즈니스 혁신을 이끄는 IT 솔루션 기술 파트너입니다'",
                "▶ [엔터프라이즈 신뢰 지표]:",
                "  · 99.98% 시스템 가동률 및 무결점 서비스 운영   · 50+ 국내외 엔터프라이즈 프로젝트 성공 레퍼런스"
            ])

    # ----------------------------------------------------
    # Slide 18: 서브페이지 - ABOUT US (4대 BIZ DOMAIN) - Three.js 삭제!
    # ----------------------------------------------------
    s18 = prs.slides[17]
    for shape in s18.shapes:
        if shape.name == "직사각형 9":
            set_shape_text(shape, [
                "01. WEB & APP DEVELOPMENT\n",
                "반응형 웹 / 모바일 앱 구축",
                "· Vue.js, React, Next.js 초고속 웹 환경",
                "· Flutter, React Native 크로스플랫폼 앱",
                "· 반응형 인터랙티브 웹 UI/UX 구축"
            ])
        elif shape.name == "TextBox 14":
            set_shape_text(shape, [
                "▶ [4대 핵심 BIZ DOMAIN 설계 및 연동]:",
                "  · Web/App, Smart Platform/ERP, Enterprise DX, Cloud Consulting",
                "  · 2x2 반응형 카드 그리드 & 감각적인 모션 배경 연동\n",
                "▶ [호버 인터랙션]:",
                "  · 마우스 호버 시 1.05배 Scale 확대 및 네온 시안(#0EA5E9) 발광 효과\n",
                "▶ [엔터프라이즈 레퍼런스 연동]:",
                "  · LX 완성창 ERP 등 실제 구축 사례와 각 사업영역 간 다이렉트 링크 제공\n",
                "▶ [반응형 레이아웃]:",
                "  · 데스크톱 2x2 그리드 -> 태블릿/모바일 1열 자동 재배치"
            ])

    # ----------------------------------------------------
    # Slide 23: 기술 아키텍처 (Three.js 및 3D 삭제, 쉬운 말로 변경)
    # ----------------------------------------------------
    s23 = prs.slides[22]
    for shape in s23.shapes:
        if shape.name == "Text 21":
            set_shape_text(shape, "3. Interactive Web & Vector Motion")
        elif shape.name == "Text 22":
            set_shape_text(shape, "가볍고 빠른 모던 웹 인터랙션 및 벡터 그래픽")
        elif shape.name == "Text 23":
            set_shape_text(shape, [
                "· 무거운 3D 플러그인을 배제하고 고화질 비디오와 가벼운 SVG 벡터 모션으로 최적화",
                "· 저사양 PC와 모바일 디바이스에서도 버벅임 없이 쾌적하고 부드러운 반응성 보장"
            ])
        elif shape.name == "Text 27":
            set_shape_text(shape, "페이지 새로고침 없는 유기적 사용자 경험")
        elif shape.name == "Text 28":
            set_shape_text(shape, [
                "· GNB 4대 메뉴(Home, About Us, Works, Contact Us) 이동 시에도 화면이 깜빡이지 않고 부드럽게 전환",
                "· 풀스크린 커튼 드로어로 모바일과 데스크톱 모두에서 일관된 프리미엄 UX 제공"
            ])

    # ----------------------------------------------------
    # Slide 24: 구축 및 배포 실행 로드맵 (Three.js 삭제)
    # ----------------------------------------------------
    s24 = prs.slides[23]
    for shape in s24.shapes:
        if shape.name == "Text 20":
            set_shape_text(shape, [
                "· Hero 비디오 배경 및 인터랙티브 모션 연동",
                "· S.01~S.04 핀 스크롤 및 에너지 스트림 개발",
                "· 차세대 인터랙티브 스타일 2x3 무박스 수치 카운트업"
            ])

    prs.save(dst_path)
    print(f"Saved successfully to {dst_path}")

if __name__ == "__main__":
    src = "WINHUB_리뉴얼_combine_v2_압축수정.backup_before_edit.pptx"
    dst = "WINHUB_리뉴얼_combine_v2_압축수정_updated.pptx"
    update_presentation(src, dst)
