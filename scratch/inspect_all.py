import pptx
import sys

def inspect_slide(prs, slide_idx):
    slide = prs.slides[slide_idx]
    print(f"=== Slide {slide_idx + 1} Shapes ===")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            texts = [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
            print(f"[{s_idx}] id={shape.shape_id}, name={shape.name}, text={texts[:2]}")

if __name__ == "__main__":
    prs = pptx.Presentation("WINHUB_리뉴얼_combine_v2_압축수정.pptx")
    for idx in [1, 2, 3, 4, 5, 8, 10, 11, 12, 13, 14, 15, 16, 17, 22, 23]:
        inspect_slide(prs, idx)
