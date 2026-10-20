# -*- coding: utf-8 -*-
import pptx
from pptx.util import Pt
from pptx.dml.color import RGBColor
import os

def format_bullet_box(shape, items):
    tf = shape.text_frame
    tf.word_wrap = True
    
    para_list = []
    for header, bullets in items:
        para_list.append((header, True, 9.5))
        for b in bullets:
            para_list.append((f"  · {b}", False, 8.5))
        para_list.append(("", False, 4.0)) # small spacing
    
    if para_list and para_list[-1][0] == "":
        para_list.pop()

    for idx, (p_text, is_bold, sz) in enumerate(para_list):
        if idx < len(tf.paragraphs):
            p = tf.paragraphs[idx]
        else:
            p = tf.add_paragraph()
        p.text = p_text
        p.font.name = "Malgun Gothic"
        p.font.size = Pt(sz)
        p.font.bold = is_bold
        p.font.color.rgb = RGBColor(30, 41, 59)
        for r in p.runs:
            r.font.name = "Malgun Gothic"
            r.font.size = Pt(sz)
            r.font.bold = is_bold
            r.font.color.rgb = RGBColor(30, 41, 59)

    if len(tf.paragraphs) > len(para_list):
        for idx in range(len(para_list), len(tf.paragraphs)):
            p = tf.paragraphs[idx]
            p.text = ""
            p.font.size = Pt(1)

def set_simple_text(shape, lines, font_name="Malgun Gothic", font_size=None, is_bold=None, color_rgb=None):
    tf = shape.text_frame
    if isinstance(lines, str):
        lines = [lines]
    
    for idx, line in enumerate(lines):
        if idx < len(tf.paragraphs):
            p = tf.paragraphs[idx]
        else:
            p = tf.add_paragraph()
        p.text = line
        if font_name:
            p.font.name = font_name
        if font_size:
            p.font.size = Pt(font_size)
        if is_bold is not None:
            p.font.bold = is_bold
        if color_rgb:
            p.font.color.rgb = color_rgb
        for r in p.runs:
            if font_name:
                r.font.name = font_name
            if font_size:
                r.font.size = Pt(font_size)
            if is_bold is not None:
                r.font.bold = is_bold
            if color_rgb:
                r.font.color.rgb = color_rgb

    if len(tf.paragraphs) > len(lines):
        for idx in range(len(lines), len(tf.paragraphs)):
            p = tf.paragraphs[idx]
            p.text = ""
            p.font.size = Pt(1)


def update_presentation(src_path, dst_path):
    prs = pptx.Presentation(src_path)

    # ----------------------------------------------------
    # Slide 2: 기획서 구성 및 목차 (Slide index 1)
    # ----------------------------------------------------
    s2 = prs.slides[1]
    for shape in s2.shapes:
        if shape.name == "Text 13":
            set_simple_text(shape, [
                "· 정적 카탈로그 웹의 한계를 넘어 인터랙티브 웹으로의 전환",
                "· 고화질 비디오 배경 & 현대적 인터랙티브(엔지니어링 신뢰) 통합 모델"
            ], font_size=8.5)
        elif shape.name == "Text 18":
            set_simple_text(shape, [
                "· 풀스크린 오버레이 메뉴로 시원하고 직관적인 메뉴 탐색 환경 제공",
                "· 페이지 이동 시 화면 깜빡임 없이 부드럽게 이어지는 연속적 화면 전환",
                "· 공간 배치 혁신: 불필요한 박스 테두리를 없애고 깔끔한 헤어라인 그리드 적용"
            ], font_size=8.5)
        elif shape.name == "Text 38":
            set_simple_text(shape, [
                "· 5단계 통합 국/영문 심볼릭 카피라이팅 매트릭스",
                "· HTML5 비디오, 모던 웹 모션 라이브러리 기술 파이프라인",
                "· 4단계 구축 및 런칭 마일스톤 (Phase 1 ~ Phase 4)"
            ], font_size=8.5)

    # ----------------------------------------------------
    # Slide 3: 핵심 기획 요약 (Slide index 2)
    # ----------------------------------------------------
    s3 = prs.slides[2]
    for shape in s3.shapes:
        if shape.name == "Text 27":
            set_simple_text(shape, "시네마틱 영상 비주얼 + 엔터프라이즈 신뢰 지표", font_size=11, is_bold=True)
        elif shape.name == "Text 29":
            set_simple_text(shape, [
                "· futuristic_key.mp4 고화질 비디오 배경이 시각적 압도감 제공",
                "· 99.98% 가동률, 50+ 대기업 레퍼런스 등 구체적 수치로 설득",
                "· 강력한 B2B 프로젝트 제안 문의 전환(Conversion) 창출"
            ], font_size=8.5)

    # ----------------------------------------------------
    # Slide 4: 유기적인 페이지 이동 및 네비게이팅 시스템 (Slide index 3)
    # ----------------------------------------------------
    s4 = prs.slides[3]
    for shape in s4.shapes:
        if shape.name == "Text 4":
            set_simple_text(shape, "깜빡임 없는 매끄러운 화면 전환과 풀스크린 커튼 메뉴 시스템", font_size=10)
        elif shape.name == "Text 12":
            set_simple_text(shape, [
                "· 화면 최상단에 항상 고정되어 언제든 원하는 메뉴로 이동 가능",
                "· 스크롤 시작 시 은은한 반투명 다크 배경으로 본문과 자연스럽게 조화",
                "· 본문 콘텐츠를 가리지 않으면서 로고 / 메뉴 버튼 구조를 깔끔하게 유지",
                "· WINHUB 4대 핵심 메뉴 [Home · About Us · Works · Contact Us] 직관적 노출"
            ], font_size=8.5)
        elif shape.name == "Text 16":
            set_simple_text(shape, [
                "· 좁은 드롭다운 창 대신 화면 전체(풀스크린)를 감각적으로 덮는 시원한 오버레이 메뉴",
                "· 메뉴 텍스트: 아래에서 위로 물 흐르듯 순차적으로 나타나는 트렌디한 모션 연출",
                "· 큼직하고 세련된 폰트를 적용하여 방문자에게 명확하고 직관적인 브랜드 인상 전달",
                "· 참조 사이트: 1367studio (https://www.1367studio.com/) 풀스크린 메뉴 스타일 벤치마킹",
                "· ESC 키 감지 및 닫기 버튼으로 언제든지 편하고 빠르게 닫을 수 있는 쉬운 사용성"
            ], font_size=8.5)
        elif shape.name == "Text 19":
            set_simple_text(shape, "3. 화면 깜빡임 없는 매끄러운 페이지 이동 (화면 끊김 방지)", font_size=10, is_bold=True)
        elif shape.name == "Text 20":
            set_simple_text(shape, [
                "· 1단계: 메뉴 클릭 시 화면이 하얗게 깜빡이지 않고 부드럽게 전환 시작",
                "· 2단계: 현재 화면이 은은하게 페이드아웃되며 다음 페이지로 자연스럽게 연결",
                "· 3단계: 상단 헤더 메뉴와 배경 무드를 안정적으로 유지하며 새 콘텐츠 신속 로딩",
                "· 4단계: 새로운 페이지의 최상단 위치로 스크롤을 매끄럽게 맞추어 즉시 탐색 지원",
                "· 5단계: 신규 페이지의 인터랙티브 요소들이 유기적으로 등장하며 방문자 몰입 유도"
            ], font_size=8.5)

    # ----------------------------------------------------
    # Slide 5: 인터랙션 및 모션 시스템 (Slide index 4)
    # ----------------------------------------------------
    s5 = prs.slides[4]
    for shape in s5.shapes:
        if shape.name == "Text 24":
            set_simple_text(shape, [
                "· 화면이 아래로 밀려나지 않고 고정된 채, 스크롤에 맞춰 텍스트와 그래픽이 자연스럽게 변화",
                "· 고화질 비디오 배경, 은은한 조명 색상, SVG 신경망 그래픽이 스크롤 흐름과 완벽 결합"
            ], font_size=8.5)
        elif shape.name == "Text 27":
            set_simple_text(shape, "텍스트 순차 등장 모션 (Line Reveal)", font_size=9, is_bold=True)

    # ----------------------------------------------------
    # Slide 6: 기술 스택 요약 테이블 (Slide index 5)
    # ----------------------------------------------------
    s6 = prs.slides[5]
    for shape in s6.shapes:
        if shape.has_table:
            table = shape.table
            for row in table.rows:
                if len(row.cells) >= 3 and any(w in row.cells[1].text for w in ["Three.js", "WebGL"]):
                    row.cells[1].text = "HTML5 Video & Modern SVG"
                    row.cells[2].text = "고화질 배경 비디오, 심플 벡터 그래픽, 실시간 부드러운 웹 모션"
                    if len(row.cells) >= 4:
                        row.cells[3].text = "futuristic_key 비디오 + 세련된 AI/레이더 그래픽 연출"
                elif len(row.cells) >= 3 and "Motion Orchestration" in row.cells[0].text:
                    row.cells[2].text = "스크롤 속도에 맞추어 화면과 텍스트가 자연스럽게 변화하는 1:1 모션 연동"
                elif len(row.cells) >= 3 and "Typography Engine" in row.cells[0].text:
                    row.cells[2].text = "텍스트 문장을 행 단위로 분할하여 아래에서 위로 부드럽게 나타나는 등장 효과"

    # ----------------------------------------------------
    # Slide 9: 네비게이션 상세 스펙 (Slide index 8)
    # ----------------------------------------------------
    s9 = prs.slides[8]
    for shape in s9.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if "Clip-path Reveal" in p.text:
                    p.text = "▶ [텍스트 순차 등장]: 01 Home부터 04 Contact Us까지 텍스트가 아래에서 위로 물 흐르듯 순차적으로 나타남"
                elif "백드롭 블러(blur: 24px)" in p.text:
                    p.text = "▶ [진입 모션]: 메뉴 버튼 클릭 시 0.4초 만에 은은한 반투명 다크 배경과 함께 풀스크린 메뉴 창이 시원하게 열림"

    # ----------------------------------------------------
    # Slide 11: 메인 화면 - Hero 섹션 (Act 1, Slide index 10)
    # ----------------------------------------------------
    s11 = prs.slides[10]
    for shape in s11.shapes:
        if shape.name == "Text 40":
            format_bullet_box(shape, [
                ("▶ [고화질 배경 영상 연동]:", [
                    "assets/mp4/futuristic_key.mp4 고화질 비디오 무한 루프 재생",
                    "다크 오버레이를 적용하여 메인 텍스트의 선명한 가독성 확보"
                ]),
                ("▶ [직관적인 메인 카피]:", [
                    "'CONNECTS TO EMPOWER YOUR SUCCESS' 슬로건으로 비즈니스 도약 비전 제시",
                    "미래를 여는 디지털 키(Gateway) 모티브를 세련된 영상과 문구로 일치화"
                ]),
                ("▶ [GNB 및 화면 이동]:", [
                    "상단 고정 로고 + 햄버거 메뉴 배치 및 아래로 부드럽게 스크롤 유도"
                ])
            ])

    # ----------------------------------------------------
    # Slide 12: 메인 화면 - S.01 Tech Innovation (Act 2, Slide index 11)
    # ----------------------------------------------------
    s12 = prs.slides[11]
    for shape in s12.shapes:
        if shape.name == "Text 34":
            set_simple_text(shape, "AI 칩셋 & 신경망 그래픽", font_size=10, is_bold=True)
        elif shape.name == "Text 36":
            set_simple_text(shape, "은은한 무드 배경 효과", font_size=7)
        elif shape.name == "Text 39":
            format_bullet_box(shape, [
                ("▶ [AI 칩셋 & 신경망 비주얼]:", [
                    "심플 AI 칩셋 SVG 그래픽과 데이터 흐름 라인이 스크롤에 맞춰 은은하게 반응",
                    "가볍고 빠른 모던 2D 벡터 그래픽으로 스마트한 기술력 표현"
                ]),
                ("▶ [핵심 기술 역량 전달]:", [
                    "AI · 클라우드 · 빅데이터 아키텍처를 엔터프라이즈 환경에 최적화하여 구현",
                    "지능의 신경망 컨셉으로 대규모 IT 시스템 구축 전문성을 직관적으로 전달"
                ]),
                ("▶ [모바일/태블릿 최적화]:", [
                    "어느 기기에서나 텍스트와 그래픽이 한눈에 들어오는 가독성 중심 디자인"
                ])
            ])

    # ----------------------------------------------------
    # Slide 13: 메인 화면 - S.02 Custom Strategy (Act 3, Slide index 12)
    # ----------------------------------------------------
    s13 = prs.slides[12]
    for shape in s13.shapes:
        if shape.name == "Text 35":
            set_simple_text(shape, [
                "360° 레이더 스캐너",
                "[맞춤형 DX 로드맵 그래픽]"
            ], font_size=8)
        elif shape.name == "Text 38":
            format_bullet_box(shape, [
                ("▶ [정밀 레이더 스캔 모션]:", [
                    "360도 회전하는 레이더 그래픽으로 고객 비즈니스 환경을 정밀 분석하는 컨셉 시각화",
                    "초록빛 포인트 레이더 그래픽과 얇은 가이드 라인의 모던한 결합"
                ]),
                ("▶ [정밀 맞춤 전략 로드맵]:", [
                    "정형화된 패키지가 아닌 기업별 요구사항에 꼭 맞춘 단계별 시스템 구축 및 운영 로드맵 설계",
                    "비즈니스 파이프라인을 정밀 분석해 산업단위 최적화 솔루션 제시"
                ]),
                ("▶ [전문적인 다크 스타일]:", [
                    "군더더기 없는 미니멀한 라인 그래픽으로 신뢰감과 기술적 정밀함을 동시에 전달"
                ])
            ])

    # ----------------------------------------------------
    # Slide 14: 메인 화면 - S.03 Enterprise Impact (Act 4, Slide index 13)
    # ----------------------------------------------------
    s14 = prs.slides[13]
    for shape in s14.shapes:
        if shape.name == "Text 32":
            set_simple_text(shape, "차세대 인터랙티브 2x3 수치 그리드", font_size=8)
        elif shape.name == "Text 53":
            format_bullet_box(shape, [
                ("▶ [숫자로 입증하는 신뢰 (2x3 그리드)]:", [
                    "99.98% 가동률, 15년+ SI 노하우, 50+ 대기업 파트너 등 핵심 수치 카운트업 효과",
                    "화면 스크롤 도달 시 수치가 0에서 목표치까지 부드럽게 올라가는 동적 연출"
                ]),
                ("▶ [직관적인 정보 구조]:", [
                    "답답한 박스 테두리를 없애고 모던한 헤어라인 구분선으로 6개 핵심 지표 정렬",
                    "시각적 피로도 없이 방대한 실적과 신뢰도를 한눈에 파악 가능"
                ]),
                ("▶ [검증된 레퍼런스 강조]:", [
                    "말로만 하는 기술이 아닌 실질적인 대규모 운영 실적으로 고객의 확실한 신뢰 확보"
                ])
            ])

    # ----------------------------------------------------
    # Slide 15: 메인 화면 - S.04 Works Showcase (Act 5 전반부, Slide index 14)
    # ----------------------------------------------------
    s15 = prs.slides[14]
    for shape in s15.shapes:
        if shape.name == "TextBox 60":
            format_bullet_box(shape, [
                ("▶ [에디토리얼 실적 인덱스 테이블]:", [
                    "ABC마트, LX하우시스, AK플라자, 신세계I&C 등 대표 프로젝트를 깔끔한 목록으로 제시",
                    "마우스를 올리면 각 프로젝트의 대표 이미지와 정보가 부드럽게 나타남"
                ]),
                ("▶ [실적 상세 페이지 연계]:", [
                    "각 프로젝트 행 클릭 시 Works 상세 화면으로 자연스럽게 이동하여 심층 내용 확인",
                    "대형 제조·유통·금융·통신 엔터프라이즈 고객사의 구축 성공 사례 직관적 노출"
                ]),
                ("▶ [세련된 엔터프라이즈 디자인]:", [
                    "얇은 구분선과 시안 컬러 포인트로 높은 시인성과 세련된 완성도 확보"
                ])
            ])

    # ----------------------------------------------------
    # Slide 16: 메인 화면 - S.05 Outro CTA (Act 5 후반부, Slide index 15)
    # ----------------------------------------------------
    s16 = prs.slides[15]
    for shape in s16.shapes:
        if shape.name == "Text 38":
            format_bullet_box(shape, [
                ("▶ [명확한 문의 유도 (CTA)]:", [
                    "'READY TO EMPOWER YOUR FUTURE?' 강력한 메시지와 함께 프로젝트 제안 문의 버튼 강조",
                    "단순한 시스템 구축을 넘어 성공적인 디지털 전환(DX)을 시작하도록 적극 유도"
                ]),
                ("▶ [하단 푸터 정보 안내]:", [
                    "WINHUB 회사 소개, 서브페이지 네비게이션(About Us, Works, Contact Us)",
                    "본사 위치, 공식 대표 연락처 및 이메일 정보를 일목요연하게 배치"
                ]),
                ("▶ [5단계 스토리텔링의 완결]:", [
                    "Act 1(미래를 여는 키 비디오)부터 Act 5(제안 및 실질적 문의)까지 막힘없는 여정 완성"
                ])
            ])

    # ----------------------------------------------------
    # Slide 17: 서브페이지 - ABOUT US (히어로, Slide index 16)
    # ----------------------------------------------------
    s17 = prs.slides[16]
    for shape in s17.shapes:
        if shape.name == "TextBox 48":
            set_simple_text(shape, "서브페이지 - ABOUT US (기업 비전 & 시네마틱 영상 히어로)", font_size=11, is_bold=True)
        elif shape.name == "Text 44":
            format_bullet_box(shape, [
                ("▶ [시네마틱 영상 비주얼 연출]:", [
                    "wire_circle_network.mp4 비디오 배경과 기업 철학을 연결하여 세련된 몰입감 제공",
                    "고화질 비디오 위에 텍스트와 핵심 수치를 배치하여 첫인상에서 깊은 신뢰 형성"
                ]),
                ("▶ [기업 비전 카피라이팅]:", [
                    "'엔터프라이즈의 지속 가능한 성장을 리드하는 WINHUB의 기술 역량과 기업 철학'",
                    "비즈니스 혁신을 이끄는 전문 IT 솔루션 기술 파트너로서의 위상 정립"
                ]),
                ("▶ [엔터프라이즈 신뢰 지표]:", [
                    "99.98% 시스템 가동률 및 50+ 대기업 고객사 성공 레퍼런스 제시"
                ])
            ])

    # ----------------------------------------------------
    # Slide 18: 서브페이지 - ABOUT US (4대 BIZ DOMAIN, Slide index 17)
    # ----------------------------------------------------
    s18 = prs.slides[17]
    for shape in s18.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if "Three.js" in p.text:
                    p.text = "· 모던 웹 인터랙티브 UI/UX 디자인 적용"
                elif "2x2 반응형 카드 그리드 & Three.js" in p.text:
                    p.text = "· 2x2 반응형 카드 그리드 & 세련된 동적 배경 연동"

    # ----------------------------------------------------
    # Slide 19: 서브페이지 - WORKS (히어로, Slide index 18)
    # ----------------------------------------------------
    s19 = prs.slides[18]
    for shape in s19.shapes:
        if shape.name == "TextBox 48":
            set_simple_text(shape, "서브페이지 - WORKS (구축 실적 & 인터랙티브 쇼케이스)", font_size=11, is_bold=True)
        elif shape.name == "Text 44":
            format_bullet_box(shape, [
                ("▶ [시네마틱 영상 비주얼 연출]:", [
                    "triangle_connect.mp4 비디오 배경으로 기술적 연결성과 첨단 프로젝트 역량 표현",
                    "영상 위로 'PROVEN ENTERPRISE DELIVERABLES' 헤드라인과 핵심 실적 배치"
                ]),
                ("▶ [실적 헤드라인 및 필터 시스템]:", [
                    "대형 제조·유통·금융·통신 엔터프라이즈 고객사의 성공적인 구축 레퍼런스",
                    "ALL, E-Commerce, Smart Factory, ERP 등 분야별로 빠르게 골라보는 카테고리 필터"
                ]),
                ("▶ [인터랙티브 프로젝트 그리드]:", [
                    "각 프로젝트 카드 호버 시 세부 정보와 썸네일이 유기적으로 나타나는 고감도 연출"
                ])
            ])

    # ----------------------------------------------------
    # Slide 21: 서브페이지 - CONTACT US (Slide index 20)
    # ----------------------------------------------------
    s21 = prs.slides[20]
    for shape in s21.shapes:
        if shape.name == "TextBox 48":
            set_simple_text(shape, "서브페이지 - CONTACT US (비즈니스 문의 & 제안 채널)", font_size=11, is_bold=True)
        elif shape.name == "Text 44":
            format_bullet_box(shape, [
                ("▶ [시네마틱 영상 비주얼 연출]:", [
                    "contact_join.mp4 비디오 배경으로 파트너십과 연결(Connect)의 의미 극대화",
                    "영상과 어우러지는 모던 다크 무드로 문의 작성 시 집중도와 신뢰감 향상"
                ]),
                ("▶ [직관적인 온라인 문의 폼]:", [
                    "기업명, 담당자 정보, 프로젝트 예산 및 일정, 문의 내용을 손쉽게 입력하는 간결한 폼",
                    "실시간 입력 검증으로 입력 오류를 방지하고 편리한 접수 경험 제공"
                ]),
                ("▶ [오시는 길 및 공식 채널]:", [
                    "본사 주소, 대표 전화, 문의 공식 이메일 및 지도 안내를 한눈에 확인 가능"
                ])
            ])

    # ----------------------------------------------------
    # Slide 23: 기술 아키텍처 및 파이프라인 (Slide index 22)
    # ----------------------------------------------------
    s23 = prs.slides[22]
    for shape in s23.shapes:
        if shape.name == "Text 21":
            set_simple_text(shape, "3. Modern SVG & HTML5 Video Engine", font_size=10, is_bold=True)
        elif shape.name == "Text 22":
            set_simple_text(shape, "고화질 비디오 배경 및 정밀 벡터 그래픽", font_size=8.5)
        elif shape.name == "Text 23":
            set_simple_text(shape, [
                "· 고화질 비디오 배경(futuristic_key.mp4)과 스크롤 연동 AI 칩셋/레이더 그래픽",
                "· 가볍고 빠른 모던 2D 벡터 모션으로 저사양 기기에서도 끊김 없는 부드러움 보장"
            ], font_size=7.5)

    # ----------------------------------------------------
    # Slide 24: 구축 및 배포 실행 로드맵 (Slide index 23)
    # ----------------------------------------------------
    s24 = prs.slides[23]
    for shape in s24.shapes:
        if shape.name == "Text 20":
            set_simple_text(shape, [
                "· Hero 고화질 비디오 배경 및 상단 헤더 연동",
                "",
                "· S.01~S.04 스크롤 연동 그래픽 및 모션 개발",
                "",
                "· 차세대 인터랙티브 스타일 2x3 무박스 수치 카운트업"
            ], font_size=8)

    prs.save(dst_path)
    print(f"Successfully saved all slides update to {dst_path}")

if __name__ == "__main__":
    src = r"c:\win-renewal\WINHUB_리뉴얼_combine_v2_압축수정.backup_before_edit.pptx"
    dst = r"c:\win-renewal\WINHUB_리뉴얼_combine_v2_압축수정.pptx"
    update_presentation(src, dst)
